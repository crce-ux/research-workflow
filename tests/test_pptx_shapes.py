# -*- coding: utf-8 -*-
"""pptx_shapes 的测试。

手工搭一个最小 .pptx（就是个 zip），逐项钉住 OOXML 的读法：
单位换算、组合变换、主题色、样式引用、虚线、箭头端、表格计数。

最后一节用 `python-pptx` 做**独立交叉验证**——它是另一份实现对同一套 OOXML 的读取。
装了 `python-pptx` 才跑，没装就 skip（`doctor.py` 的常规测试不带它）。
"""
import base64
import zipfile

import pytest

import pptx_shapes as PS

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
THEME_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"

W_EMU, H_EMU = 12192000, 6858000          # 13.333in × 7.5in
K = 1600 / W_EMU                          # EMU → px


def _ns(root):
    return (f'xmlns:a="{A}" xmlns:p="{P}" xmlns:r="{R}"')


def _presentation(w=W_EMU, h=H_EMU, dts=""):
    return f"""<?xml version="1.0"?>
<p:presentation {_ns('')}>
  <p:sldIdLst><p:sldId id="256" r:id="rId2"/></p:sldIdLst>
  <p:sldSz cx="{w}" cy="{h}"/>
  {dts}
</p:presentation>"""


DTS_18 = ('<p:defaultTextStyle><a:lvl1pPr><a:defRPr sz="1800"/>'
          '</a:lvl1pPr></p:defaultTextStyle>')

PRES_RELS = f"""<?xml version="1.0"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="{THEME_REL}" Target="theme/theme1.xml"/>
  <Relationship Id="rId2"
    Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
    Target="slides/slide1.xml"/>
  <Relationship Id="rId3"
    Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster"
    Target="slideMasters/slideMaster1.xml"/>
</Relationships>"""

THEME = f"""<?xml version="1.0"?>
<a:theme {_ns('')} name="T">
  <a:themeElements><a:clrScheme name="T">
    <a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1>
    <a:lt1><a:sysClr val="window" lastClr="FFFFFF"/></a:lt1>
    <a:dk2><a:srgbClr val="44546A"/></a:dk2>
    <a:lt2><a:srgbClr val="E7E6E6"/></a:lt2>
    <a:accent1><a:srgbClr val="0046A4"/></a:accent1>
    <a:accent2><a:srgbClr val="C00000"/></a:accent2>
  </a:clrScheme></a:themeElements>
</a:theme>"""


def _slide(body):
    return f"""<?xml version="1.0"?>
<p:sld {_ns('')}><p:cSld><p:spTree>
  <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
  <p:grpSpPr/>
  {body}
</p:spTree></p:cSld></p:sld>"""


def _sp(sid, name, frame, spPr_extra="", tx="", sz=None, color=None, style=""):
    x, y, cx, cy = frame
    run = (f'<a:r><a:rPr sz="{sz}"><a:solidFill><a:srgbClr val="{color}"/>'
           f'</a:solidFill></a:rPr><a:t>{tx}</a:t></a:r>') if tx else ""
    return f"""<p:sp>
  <p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>
    {spPr_extra}
  </p:spPr>
  {style}
  <p:txBody><a:bodyPr/><a:lstStyle/><a:p>{run}</a:p></p:txBody>
</p:sp>"""


EMPTY_RELS = ('<?xml version="1.0"?><Relationships '
              'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>')


def _build(tmp_path, body, extra_parts=(), master=None, layout=None,
           slide_rels=None, canvas=(W_EMU, H_EMU), dts=""):
    p = tmp_path / "t.pptx"
    parts = {
        "ppt/presentation.xml": _presentation(*canvas, dts=dts),
        "ppt/_rels/presentation.xml.rels": PRES_RELS,
        "ppt/theme/theme1.xml": THEME,
        "ppt/slides/slide1.xml": _slide(body),
        "ppt/slides/_rels/slide1.xml.rels": slide_rels or EMPTY_RELS,
    }
    if master:
        parts["ppt/slideMasters/slideMaster1.xml"] = master
    if layout:
        parts["ppt/slideLayouts/slideLayout1.xml"] = layout
    for k, v in extra_parts:
        parts[k] = v
    with zipfile.ZipFile(p, "w") as zf:
        for k, v in parts.items():
            zf.writestr(k, v)
    return p


RECT = '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'


def _read(tmp_path, body, **kw):
    shapes, stats = PS.read_pptx(_build(tmp_path, body, **kw), deck="d")
    return {s["sid"]: s for s in shapes}, stats, shapes


# ══════════════════════════════════════════════════════════
#  单位与几何
# ══════════════════════════════════════════════════════════
def test_emu_is_scaled_to_the_library_canvas(tmp_path):
    by, _, _ = _read(tmp_path, _sp(2, "块", (W_EMU // 10, H_EMU // 10,
                                            W_EMU // 2, H_EMU // 2), RECT))
    s = by["2"]
    assert (round(s["x"], 1), round(s["y"], 1)) == (160.0, 90.0)
    assert (round(s["w"], 1), round(s["h"], 1)) == (800.0, 450.0)


def test_point_size_and_line_width(tmp_path):
    """`a:rPr/@sz` 是 1/100 pt；`a:ln/@w` 是 EMU。折算后 1 pt = 1.6667 px。"""
    ln = '<a:ln w="12700"><a:solidFill><a:srgbClr val="C00000"/></a:solidFill></a:ln>'
    by, _, _ = _read(tmp_path, _sp(2, "块", (0, 0, 1000, 1000), RECT + ln,
                                   tx="题", sz=2400, color="FFFFFF"))
    s = by["2"]
    assert round(s["size_px"], 2) == 40.0            # 24 pt
    assert round(s["stroke_w"], 2) == 1.67           # 1 pt
    assert s["stroke"] == "#C00000"


def test_group_transform_is_applied(tmp_path):
    """组合的子形状坐标在组自己的坐标系里，要按 chOff / chExt 换算。

    组的外框 12192000×6858000，子坐标系 24384000×13716000 → 缩放 0.5。
    子形状 off=(2438400,1371600) ext=(2438400,1371600) → 换算后 (1219200,685800,1219200,685800)。
    """
    body = f"""<p:grpSp>
  <p:nvGrpSpPr><p:cNvPr id="9" name="组合"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
  <p:grpSpPr><a:xfrm>
    <a:off x="0" y="0"/><a:ext cx="12192000" cy="6858000"/>
    <a:chOff x="0" y="0"/><a:chExt cx="24384000" cy="13716000"/>
  </a:xfrm></p:grpSpPr>
  {_sp(10, "子块", (2438400, 1371600, 2438400, 1371600), RECT)}
</p:grpSp>"""
    by, _, _ = _read(tmp_path, body)
    kid = by["10"]
    assert (round(kid["x"], 1), round(kid["y"], 1)) == (160.0, 90.0)
    assert (round(kid["w"], 1), round(kid["h"], 1)) == (160.0, 90.0)
    assert kid["from_group"] is True
    assert by["9"]["kind"] == "group"


# ══════════════════════════════════════════════════════════
#  颜色
# ══════════════════════════════════════════════════════════
def test_scheme_colour_is_resolved_from_the_theme(tmp_path):
    fill = '<a:solidFill><a:schemeClr val="accent1"/></a:solidFill>'
    by, _, _ = _read(tmp_path, _sp(2, "块", (0, 0, 1000, 1000), RECT + fill))
    assert by["2"]["fill"] == "#0046A4"


def test_style_fill_ref_is_used_when_there_is_no_explicit_fill(tmp_path):
    """多数 rect 的颜色来自 `p:style` 的 fillRef。只读显式填充会全落成空，
    这些形状随后会被 A6 虚线框判据误收。"""
    style = ('<p:style><a:lnRef idx="1"><a:schemeClr val="accent2"/></a:lnRef>'
             '<a:fillRef idx="2"><a:schemeClr val="accent1"/></a:fillRef>'
             '<a:fontRef idx="minor"/></p:style>')
    by, _, _ = _read(tmp_path, _sp(2, "块", (0, 0, 1000, 1000), RECT, style=style))
    assert by["2"]["fill"] == "#0046A4"
    assert by["2"]["stroke"] == "#C00000"
    assert "from-style-ref" in by["2"]["notes"]


def test_no_fill_is_not_unresolved(tmp_path):
    by, _, _ = _read(tmp_path, _sp(2, "块", (0, 0, 1000, 1000),
                                   RECT + "<a:noFill/>"))
    assert by["2"]["fill"] == "none"


def test_theme_unresolved_is_flagged(tmp_path):
    fill = '<a:solidFill><a:schemeClr val="accent9"/></a:solidFill>'
    by, _, _ = _read(tmp_path, _sp(2, "块", (0, 0, 1000, 1000), RECT + fill))
    assert by["2"]["fill"] == "unresolved"
    assert "theme-unresolved" in by["2"]["notes"]


# ══════════════════════════════════════════════════════════
#  形状族与连接件
# ══════════════════════════════════════════════════════════
def test_prst_and_round_rect_adjustment(tmp_path):
    geom = ('<a:prstGeom prst="roundRect"><a:avLst>'
            '<a:gd name="adj" fmla="val 50000"/></a:avLst></a:prstGeom>')
    by, _, _ = _read(tmp_path, _sp(2, "胶囊", (0, 0, 1219200, 6096000), geom))
    assert by["2"]["prst"] == "roundRect"
    assert by["2"]["av_adj"] == 0.5
    assert round(by["2"]["rx"], 1) == round(0.5 * 160.0, 1)


def test_dash_and_arrowhead_are_read_directly(tmp_path):
    """`a:prstDash` 与 `a:tailEnd` 直接读，不用从 SVG 的 marker 反推。"""
    ln = ('<a:ln w="25400"><a:solidFill><a:srgbClr val="0046A4"/></a:solidFill>'
          '<a:prstDash val="dash"/><a:tailEnd type="triangle"/></a:ln>')
    by, _, _ = _read(tmp_path, _sp(2, "线", (0, 0, 1000, 0),
                                   '<a:prstGeom prst="line"><a:avLst/></a:prstGeom>' + ln))
    assert by["2"]["dash"] == "dash"
    assert by["2"]["head"] == "end"


def test_graphic_frame_with_a_table_is_unsupported(tmp_path):
    body = f"""<p:graphicFrame>
  <p:nvGraphicFramePr><p:cNvPr id="7" name="表格 1"/><p:cNvGraphicFramePr/>
    <p:nvPr/></p:nvGraphicFramePr>
  <p:xfrm><a:off x="0" y="0"/><a:ext cx="1000000" cy="1000000"/></p:xfrm>
  <a:graphic><a:graphicData uri="…/table"><a:tbl><a:tr><a:tc>
    <a:txBody><a:p><a:r><a:t>合计</a:t></a:r></a:p></a:txBody>
  </a:tc></a:tr></a:tbl></a:graphicData></a:graphic>
</p:graphicFrame>"""
    _, stats, shapes = _read(tmp_path, body)
    gf = [s for s in shapes if s["kind"] == "graphic-frame"][0]
    assert gf["unsupported"] is True
    assert "tbl" in gf["notes"]
    assert gf["cells"] == [["合计"]]
    assert stats["tbl"] == 1


def test_graphic_frame_frame_is_read(tmp_path):
    """`p:graphicFrame` 的帧**就是** `p:xfrm` 本身，不是它的子元素。

    旧写法找到 `p:xfrm` 后又去里面找 `a:xfrm`，找不到返回 None —— 结果全册
    22 处表格与图表的坐标全落成 `0,0 0×0`，主内容是表格的页连行列都算不出来。
    """
    body = """<p:graphicFrame>
  <p:nvGraphicFramePr><p:cNvPr id="7" name="表格 1"/><p:cNvGraphicFramePr/>
    <p:nvPr/></p:nvGraphicFramePr>
  <p:xfrm><a:off x="1219200" y="685800"/><a:ext cx="6096000" cy="1371600"/></p:xfrm>
  <a:graphic><a:graphicData uri="…/table"><a:tbl><a:tr><a:tc>
    <a:txBody><a:p><a:r><a:t>合计</a:t></a:r></a:p></a:txBody>
  </a:tc></a:tr></a:tbl></a:graphicData></a:graphic>
</p:graphicFrame>"""
    _, stats, shapes = _read(tmp_path, body)
    gf = [s for s in shapes if s["kind"] == "graphic-frame"][0]
    assert (round(gf["x"], 1), round(gf["y"], 1)) == (160.0, 90.0)
    assert (round(gf["w"], 1), round(gf["h"], 1)) == (800.0, 180.0)
    assert stats["no-frame"] == 0


# ══════════════════════════════════════════════════════════
#  母版与版式
# ══════════════════════════════════════════════════════════
def test_master_shapes_get_their_own_scope(tmp_path):
    """母版与版式形状另走一遍并标 scope——不求占位符继承，与 ppt-master 的做法相同。"""
    master = f"""<?xml version="1.0"?>
<p:sldMaster {_ns('')}><p:cSld><p:spTree>
  <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
  <p:grpSpPr/>
  {_sp(20, "母版条", (0, 0, 1000, 1000), RECT)}
</p:spTree></p:cSld></p:sldMaster>"""
    by, _, _ = _read(tmp_path, _sp(2, "内容", (0, 0, 1000, 1000), RECT), master=master)
    assert by["20"]["scope"] == "master"
    assert by["2"]["scope"] == "slide"


def test_page_numbers_start_at_one_and_are_sequential(tmp_path):
    body = "".join(_sp(i, f"块{i}", (0, 0, 1000, 1000), RECT) for i in range(2, 5))
    _, _, shapes = _read(tmp_path, body)
    assert {s["page"] for s in shapes} == {1}


# ══════════════════════════════════════════════════════════
#  非 16:9
# ══════════════════════════════════════════════════════════
def test_non_16x9_scales_each_axis_separately(tmp_path):
    """4:3 的模板：x 按宽缩放、y 按高缩放。

    形状在画布上的**相对位置与占比**保持不变，绝对宽高比会变。
    占比类判据（横色带 `h ≤ 0.15H`、竖标签 `w ≤ 0.12W`）靠这个才成立。
    """
    w43, h43 = 9144000, 6858000            # 10in × 7.5in
    p = _build(tmp_path, _sp(2, "块", (w43 // 4, h43 // 4, w43 // 2, h43 // 2), RECT),
               canvas=(w43, h43))
    shapes, stats = PS.read_pptx(p, deck="d")
    s = next(x for x in shapes if x["scope"] == "slide")
    assert (round(s["x"], 1), round(s["y"], 1)) == (400.0, 225.0)
    assert (round(s["w"], 1), round(s["h"], 1)) == (800.0, 450.0)
    assert stats["non-16x9"] == 1


# ══════════════════════════════════════════════════════════
#  占位符帧继承
# ══════════════════════════════════════════════════════════
SLIDE_LAYOUT_REL = ('http://schemas.openxmlformats.org/officeDocument/2006/relationships'
                    '/slideLayout')
MASTER_REL = ('http://schemas.openxmlformats.org/officeDocument/2006/relationships'
              '/slideMaster')


def _rels(*pairs):
    body = "".join(
        f'<Relationship Id="{rid}" Type="{typ}" Target="{tgt}"/>' for rid, typ, tgt in pairs)
    return ('<?xml version="1.0"?><Relationships '
            'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f"{body}</Relationships>")


def test_placeholder_frame_is_inherited_from_the_layout(tmp_path):
    """占位符自己没有 `a:xfrm` 时，帧写在版式里。

    不补这一步，模板里的标题与正文就全是零坐标——A1 色带、B 分栏都算不出来。
    """
    layout = f"""<?xml version="1.0"?>
<p:sldLayout {_ns('')}><p:cSld><p:spTree>
  <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
  <p:grpSpPr/>
  <p:sp><p:nvSpPr><p:cNvPr id="2" name="Title"/><p:cNvSpPr/>
    <p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr>
    <p:spPr><a:xfrm><a:off x="1219200" y="685800"/>
      <a:ext cx="6096000" cy="1371600"/></a:xfrm>{RECT}</p:spPr>
    <p:txBody><a:bodyPr/><a:p/></p:txBody>
  </p:sp>
</p:spTree></p:cSld></p:sldLayout>"""

    ph_sp = """<p:sp>
  <p:nvSpPr><p:cNvPr id="5" name="标题占位"/><p:cNvSpPr/>
    <p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr>
  <p:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
  <p:txBody><a:bodyPr/><a:p><a:r><a:rPr sz="2400"/><a:t>题</a:t></a:r></a:p></p:txBody>
</p:sp>"""

    p = _build(tmp_path, ph_sp, layout=layout,
               slide_rels=_rels(("rId1", SLIDE_LAYOUT_REL, "../slideLayouts/slideLayout1.xml")))
    shapes, stats = PS.read_pptx(p, deck="d")
    got = next(s for s in shapes if s["scope"] == "slide")
    assert (round(got["x"], 1), round(got["y"], 1)) == (160.0, 90.0)
    assert (round(got["w"], 1), round(got["h"], 1)) == (800.0, 180.0)
    assert "frame-from-layout" in got["notes"]
    assert "no-frame" not in got["notes"]


def test_slide_records_which_layout_it_uses(tmp_path):
    """幻灯片要记下自己用的是哪个版式——版式上的 logo 靠这个值挑出来。

    不记的话「这一页的版式装饰」只能退化成「全册所有版式的装饰」，
    一份 11 个版式的模板会把 11 个 logo 全叠上去。
    """
    layout = f"""<?xml version="1.0"?>
<p:sldLayout {_ns('')}><p:cSld><p:spTree>
  <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
  <p:grpSpPr/>
  {_sp(20, "版式装饰", (0, 0, 1000, 1000), RECT)}
</p:spTree></p:cSld></p:sldLayout>"""
    p = _build(tmp_path, _sp(2, "内容", (0, 0, 1000, 1000), RECT), layout=layout,
               slide_rels=_rels(("rId1", SLIDE_LAYOUT_REL,
                                 "../slideLayouts/slideLayout1.xml")))
    shapes, _ = PS.read_pptx(p, deck="d")
    slide = next(s for s in shapes if s["scope"] == "slide")
    lay = next(s for s in shapes if s["scope"] == "layout")
    assert slide["layout_part"] == "ppt/slideLayouts/slideLayout1.xml"
    assert lay["part"] == "ppt/slideLayouts/slideLayout1.xml"
    assert slide["placeholder"] is False


def test_a_placeholder_is_flagged(tmp_path):
    """版式上大量「单击此处编辑母版标题样式」的提示框要标出来。

    不标的话它们会被当成装饰收进模块，重放时页面上多出一排提示文字。
    """
    ph_sp = """<p:sp>
  <p:nvSpPr><p:cNvPr id="5" name="标题占位"/><p:cNvSpPr/>
    <p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr>
  <p:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
  <p:txBody><a:bodyPr/><a:p/></p:txBody>
</p:sp>"""
    _, _, shapes = _read(tmp_path, ph_sp)
    assert next(s for s in shapes if s["scope"] == "slide")["placeholder"] is True


def test_placeholder_without_any_frame_stays_flagged(tmp_path):
    """版式里也找不到帧，就照实标 `no-frame`，不要编一个坐标出来。"""
    ph_sp = """<p:sp>
  <p:nvSpPr><p:cNvPr id="5" name="孤儿占位"/><p:cNvSpPr/>
    <p:nvPr><p:ph type="body"/></p:nvPr></p:nvSpPr>
  <p:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
  <p:txBody><a:bodyPr/><a:p/></p:txBody>
</p:sp>"""
    _, stats, shapes = _read(tmp_path, ph_sp)
    got = next(s for s in shapes if s["scope"] == "slide")
    assert "no-frame" in got["notes"]
    assert stats["no-frame"] == 1


# ══════════════════════════════════════════════════════════
#  与 python-pptx 的独立交叉验证
# ══════════════════════════════════════════════════════════
def test_cross_check_against_python_pptx(tmp_path):
    """`python-pptx` 是另一份实现对同一套 OOXML 的读取。

    两边逐形状比几何与文字。对得上，说明本文件读对了。
    装了 `python-pptx` 才跑：`uv run --with pytest --with python-pptx python -m pytest`
    """
    pptx = pytest.importorskip("pptx")
    from pptx.util import Emu

    prs = pptx.Presentation()
    prs.slide_width, prs.slide_height = Emu(W_EMU), Emu(H_EMU)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(1219200), Emu(685800), Emu(6096000), Emu(1371600))
    box.text_frame.text = "交叉验证"
    path = tmp_path / "x.pptx"
    prs.save(str(path))

    mine, _ = PS.read_pptx(path, deck="x")
    # ⚠️ 幻灯片 / 版式 / 母版**各自从 1 开始编号**，`sid` 只在同一部件内唯一。
    #    只按 sid 建索引会把三个部件的形状混在一起。
    mine = {s["sid"]: s for s in mine if s["scope"] == "slide"}
    for sh in slide.shapes:
        got = mine.get(str(sh.shape_id))
        assert got is not None, f"漏了形状 {sh.shape_id}"
        for axis, emu in (("x", sh.left), ("y", sh.top), ("w", sh.width), ("h", sh.height)):
            assert abs(got[axis] - emu * K) < 0.5, f"{axis} 对不上：{got[axis]} vs {emu * K}"
        assert (got["text"] or "").strip() == (sh.text_frame.text or "").strip()


# ══════════════════════════════════════════════════════════
#  自定义几何 · 绑定配色 · 渐变
# ══════════════════════════════════════════════════════════
CUST = ('<a:custGeom><a:avLst/><a:gdLst/>'
        '<a:pathLst><a:path w="100" h="50">'
        '<a:moveTo><a:pt x="0" y="0"/></a:moveTo>'
        '<a:lnTo><a:pt x="100" y="0"/></a:lnTo>'
        '<a:cubicBezTo><a:pt x="100" y="25"/><a:pt x="50" y="50"/>'
        '<a:pt x="0" y="50"/></a:cubicBezTo>'
        '<a:close/></a:path></a:pathLst></a:custGeom>')


def test_custom_geometry_becomes_path_data(tmp_path):
    """`a:custGeom` 占非文字内容件的 53%——图标、箭头、手绘形全在这一类。

    不读的话这些形状的 `prst` 是空的，重放时全被跳过。
    """
    by, _, _ = _read(tmp_path, _sp(2, "图标", (0, 0, 1000, 1000), CUST))
    p = by["2"]["path"]
    assert p["w"] == 100 and p["h"] == 50
    assert p["d"].startswith("M0.0,0.0 L100.0,0.0 C")
    assert p["d"].rstrip().endswith("Z")


def test_a_path_without_w_falls_back_to_the_shape_extent(tmp_path):
    """`a:path/@w` 与 `@h` 可以一个都不给，缺省时取形状自己的 extent。

    ⚠️ 不兜这一下，「任意多边形」那类只给 `h` 的直条会被整件丢掉——
    实测学术答辩-01 有三件（都是不带 `@w` 的竖线）。
    """
    nogeom = ('<a:custGeom><a:avLst/><a:gdLst/>'
              '<a:pathLst><a:path h="1193800">'
              '<a:moveTo><a:pt x="0" y="0"/></a:moveTo>'
              '<a:lnTo><a:pt x="0" y="1193800"/></a:lnTo>'
              '</a:path></a:pathLst></a:custGeom>')
    by, _, _ = _read(tmp_path, _sp(2, "任意多边形", (0, 0, 46076, 1374373), nogeom))
    p = by["2"]["path"]
    assert p is not None and p["w"] == 46076      # 缺省 = ext/@cx
    assert p["h"] == 1193800


def test_lum_mod_and_lum_off_produce_a_tint(tmp_path):
    """`lumMod` / `lumOff` 把主题色拉出一族浓淡。6 份模板实测 1444 / 546 个。

    ⚠️ 不读的话每一个浓淡都当成基色。浙大蓝那份结论页的浅蓝底是
       `lumMod 20% + lumOff 80%`，不读就记成深蓝 `#003F88`，
       统一配色后整块变成饱和蓝，块上的字读不了。
    """
    body = _sp(2, "浅底块", (0, 0, 1000, 1000),
               RECT + '<a:solidFill><a:schemeClr val="accent1">'
                      '<a:lumMod val="20000"/><a:lumOff val="80000"/>'
                      '</a:schemeClr></a:solidFill>')
    by, _, _ = _read(tmp_path, body)
    # accent1 = 0046A4，L=0.32 → L*0.2+0.8 = 0.864，落成浅蓝
    assert by["2"]["fill"] != "#0046A4"
    h, s, l = PS._hsl_of(by["2"]["fill"])
    assert l > 0.8


def test_shade_darkens(tmp_path):
    """`shade` 压暗，与 `lumMod` 同一支。6 份模板实测 963 个。"""
    body = _sp(2, "深条", (0, 0, 1000, 1000),
               RECT + '<a:solidFill><a:schemeClr val="accent1">'
                      '<a:shade val="50000"/></a:schemeClr></a:solidFill>')
    by, _, _ = _read(tmp_path, body)
    h, s, l = PS._hsl_of(by["2"]["fill"])
    assert l < 0.32


def test_a_semi_transparent_fill_keeps_its_alpha(tmp_path):
    """`a:alpha` 要读。7 份模板实测 821 处。

    ⚠️ 不读的话 4% 透明度的浅底会画成饱和实色——
       国自然6红色 p11 那三张卡片就是这么从「白底红字」变成「红底红字」的。
    """
    body = _sp(2, "浅底卡", (0, 0, 1000, 1000),
               RECT + '<a:solidFill><a:srgbClr val="C00000">'
                      '<a:alpha val="4000"/></a:srgbClr></a:solidFill>')
    by, _, _ = _read(tmp_path, body)
    assert by["2"]["fill"] == "#C00000"
    assert by["2"]["fill_alpha"] == 0.04


def test_an_opaque_fill_reports_full_alpha(tmp_path):
    by, _, _ = _read(tmp_path, _sp(2, "实底", (0, 0, 1000, 1000),
                                   RECT + '<a:solidFill><a:srgbClr val="C00000"/>'
                                          '</a:solidFill>'))
    assert by["2"]["fill_alpha"] == 1.0


def test_bound_scheme_colours_resolve_through_the_colour_map(tmp_path):
    """`bg1` / `tx1` 不是色板槽位，要走母版 `p:clrMap` 折一层。

    学术答辩-01 那 160 个图标填的就是 `bg1`；不解的话画出来是空的。
    """
    master = f"""<?xml version="1.0"?>
<p:sldMaster {_ns('')}>
  <p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2"/>
</p:sldMaster>"""
    body = _sp(2, "图标", (0, 0, 1000, 1000),
               CUST + '<a:solidFill><a:schemeClr val="bg1"/></a:solidFill>')
    by, _, _ = _read(tmp_path, body, master=master)
    assert by["2"]["fill"] == "#FFFFFF"          # bg1 → lt1 → window → 白


def test_gradient_stops_are_read(tmp_path):
    """渐变占非文字内容件的 8.4%。不记首末色的话深色渐变底画不出来。"""
    gf = ('<a:gradFill><a:gsLst>'
          '<a:gs pos="0"><a:srgbClr val="0B2B6B"/></a:gs>'
          '<a:gs pos="100000"><a:srgbClr val="11754F"/></a:gs>'
          '</a:gsLst><a:lin ang="5400000"/></a:gradFill>')
    by, _, _ = _read(tmp_path, _sp(2, "底", (0, 0, 1000, 1000), RECT + gf))
    g = by["2"]["grad"]
    assert g["stops"] == [[0.0, "#0B2B6B"], [1.0, "#11754F"]]
    assert g["ang"] == 90.0                      # 5400000 / 60000


# ══════════════════════════════════════════════════════════
#  字号退路
# ══════════════════════════════════════════════════════════
def _sp_raw(sid, frame, inner):
    """直接给 `p:txBody` 的内容，用来构造「没有显式字号」的各种写法。"""
    x, y, cx, cy = frame
    return f"""<p:sp>
  <p:nvSpPr><p:cNvPr id="{sid}" name="s"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>
    {RECT}
  </p:spPr>
  <p:txBody><a:bodyPr/>{inner}</p:txBody>
</p:sp>"""


def test_size_falls_back_to_end_para_rpr(tmp_path):
    """run 没写字号、`endParaRPr` 写了——按 `endParaRPr` 取。"""
    inner = ('<a:lstStyle/><a:p><a:r><a:t>题</a:t></a:r>'
             '<a:endParaRPr sz="1400"/></a:p>')
    by, _, _ = _read(tmp_path, _sp_raw(2, (0, 0, 1000, 1000), inner), dts=DTS_18)
    s = by["2"]
    assert round(s["size_px"], 2) == round(14 * 1.6667, 2)     # 14 pt，不是 deck 的 18
    assert "size-endpara" in s["notes"]


def test_size_falls_back_to_the_deck_default_text_style(tmp_path):
    """形状里一处字号都没有——退到 `p:defaultTextStyle`，并把来源记进 notes。

    实测（2026-10-15）：答辩模板里 25–31% 的文字形状真的没有字号，
    不是读取器漏读。四份 deck 的 `p:defaultTextStyle` 都是 18 pt，
    与封面「导  师：XXX」实测的 30 px（=18 pt）对得上。
    """
    inner = '<a:lstStyle/><a:p><a:r><a:t>题</a:t></a:r></a:p>'
    by, _, _ = _read(tmp_path, _sp_raw(2, (0, 0, 1000, 1000), inner), dts=DTS_18)
    s = by["2"]
    assert round(s["size_px"], 2) == 30.0                      # 18 pt
    assert "size-deck-default" in s["notes"]


def test_no_source_at_all_leaves_the_size_unset(tmp_path):
    """连 deck 默认值都没有时不猜数字——留 0，也不加来源码。"""
    inner = '<a:lstStyle/><a:p><a:r><a:t>题</a:t></a:r></a:p>'
    by, _, _ = _read(tmp_path, _sp_raw(2, (0, 0, 1000, 1000), inner))
    s = by["2"]
    assert s["size_px"] == 0.0
    assert not [n for n in s["notes"] if n.startswith("size-")]


def test_empty_lvl_style_is_not_mistaken_for_the_next_level(tmp_path):
    """空的元素在 ElementTree 里是假值。

    退路要按段落自己的层级取（`pPr/@lvl="1"` → `lvl2pPr`）。
    写成 `find(lvl2pPr) or find(lvl1pPr)` 时，空的 `lvl2pPr` 会被判成「不存在」，
    于是错误地退到 `lvl1pPr` 的 3000——属于别的层级。
    """
    inner = ('<a:lstStyle><a:lvl1pPr><a:defRPr sz="3000"/></a:lvl1pPr>'
             '<a:lvl2pPr/></a:lstStyle>'
             '<a:p><a:pPr lvl="1"/><a:r><a:t>题</a:t></a:r></a:p>')
    by, _, _ = _read(tmp_path, _sp_raw(2, (0, 0, 1000, 1000), inner), dts=DTS_18)
    s = by["2"]
    assert round(s["size_px"], 2) == 30.0                      # 本级的 lstStyle 是空的
    assert round(s["size_px"], 2) != 50.0                      # 不是上一级的 30 pt
